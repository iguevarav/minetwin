import base64
import re
from dataclasses import replace
from io import BytesIO

import matplotlib
import numpy as np
import plotly.graph_objects as go
import pytest
from openpyxl import load_workbook
from PIL import Image
from streamlit.testing.v1 import AppTest

from minetwin.explanations import describe
from minetwin.exports import (
    DPI,
    explained_rows,
    matplotlib_jpg,
    plotly_jpg,
    table_xlsx,
)
from minetwin.learning import legacy_report
from minetwin.learning.article_report import _html_report, _markdown
from minetwin.study import _html_table

matplotlib.use("Agg")


def test_excel_preserves_values_missing_cells_and_literal_text():
    rows = [
        {"Sensor": "171_0", "Valor": np.float64(1.25), "Ausente": None},
        {"Sensor": "=1+1", "Valor": np.int64(3), "Ausente": float("nan")},
        {"Sensor": "áéη", "Valor": None, "Ausente": True},
    ]
    notes = explained_rows(describe("telemetry"), rows)
    book = load_workbook(BytesIO(table_xlsx(rows, notes)))
    assert list(book["Datos"].values) == [
        ("Sensor", "Valor", "Ausente"),
        ("171_0", 1.25, None),
        ("=1+1", 3, None),
        ("áéη", None, True),
    ]
    assert book["Datos"]["A3"].data_type == "s"
    assert list(book["Interpretación"].values) == notes.items()


def test_csv_export_keeps_metrics_numeric_and_identifiers_literal():
    rows = [
        {
            "vehicle_id": "123",
            "mean_cost": "0.25",
            "examples": "50",
            "selected": "True",
            "code": "0012",
        }
    ]
    book = load_workbook(
        BytesIO(table_xlsx(rows, describe("performance"), parse_numbers=True))
    )
    assert list(book["Datos"].values)[1] == ("123", 0.25, 50, True, "0012")


@pytest.mark.parametrize("renderer", ["plotly", "matplotlib"])
def test_image_is_rendered_at_600_dpi_and_contains_explanation(renderer):
    notes = replace(describe("performance"), result="Valor observado: 0.75")
    if renderer == "plotly":
        figure = go.Figure(go.Bar(x=["Régimen A"], y=[0.75]))
        original = figure.to_json()
        content = plotly_jpg(figure, notes)
        assert figure.to_json() == original
        expected_width, chart_height = 4800, 3000
    else:
        import matplotlib.pyplot as plt

        figure, axis = plt.subplots(figsize=(4, 2))
        axis.bar(["Régimen A"], [0.75])
        try:
            content = matplotlib_jpg(figure, notes)
        finally:
            plt.close(figure)
        expected_width, chart_height = 2400, 1200
    with Image.open(BytesIO(content)) as image:
        assert image.format == "JPEG"
        assert image.info["dpi"] == (DPI, DPI)
        assert image.width == expected_width
        assert image.height > chart_height
        # Explanatory footer has dark text, rather than just extra blank pixels.
        footer = np.asarray(image.crop((0, chart_height, image.width, image.height)))
        assert np.count_nonzero(footer < 100) > 1000


def test_html_report_embeds_a_real_excel_with_same_columns_and_explanation():
    section = _html_table(
        "Calidad", [{"count": 3, "extra": 99}], ("count",), "Mayor es mejor."
    )
    encoded = re.search(r"base64,([^\"]+)", section).group(1)
    book = load_workbook(BytesIO(base64.b64decode(encoded)))
    assert list(book["Datos"].values) == [("count",), (3,)]
    assert "Mayor es mejor." in section
    assert "Cómo se obtiene" in section
    assert "3" in book["Interpretación"]["B5"].value


def test_article_reports_link_excel_and_jpg_and_explain_each_result(tmp_path):
    tables = {"eda_clases.csv": [{"partition": "train", "class": 0, "examples": 3}]}
    figures = [str(tmp_path / "distribucion_clases.jpg")]
    _markdown(tmp_path / "resultados.md", tables, {}, figures)
    markdown = (tmp_path / "resultados.md").read_text(encoding="utf-8")
    html = _html_report(tables, figures)
    for report in (markdown, html):
        assert "eda_clases.xlsx" in report
        assert "distribucion_clases.jpg" in report
        assert "300 dpi" not in report
        assert report.count("Cómo se obtiene") == 2


def test_screen_exports_are_lazy_and_update_when_data_changes(monkeypatch):
    def unexpected(*args, **kwargs):
        raise AssertionError("Rendering a view must not generate a download eagerly")

    import minetwin.presentation_ui as ui

    monkeypatch.setattr(ui, "table_xlsx", unexpected)
    monkeypatch.setattr(ui, "_jpg", unexpected)
    app = AppTest.from_string("""
import streamlit as st
import plotly.graph_objects as go
from minetwin.explanations import describe
from minetwin.presentation_ui import render_table, render_plot
value = st.slider("Valor", 0, 10, 3)
render_table([{"Valor": value}], explanation=describe("telemetry"), key="live")
render_plot(go.Figure(go.Bar(x=["Caso"], y=[value])), explanation=describe("performance"), key="live")
""").run()
    assert not app.exception
    assert len(app.get("download_button")) == 2
    assert any("máximo 3" in caption.value for caption in app.caption)
    app.slider[0].set_value(8).run()
    assert not app.exception
    assert any("máximo 8" in caption.value for caption in app.caption)


def test_legacy_upgrade_preserves_narrative_and_replaces_only_its_figures(
    tmp_path, monkeypatch
):
    import csv

    fixtures = {
        "train_validation_class_distribution": [
            {"split": "train", "class": 0, "examples": 10}
        ],
        "variant_summary": [
            {"variant": "base", "mean_cost_mean": 0.5, "mean_cost_sd": 0.1}
        ],
        "baseline_summary": [
            {"model": "rule", "mean_cost_mean": 1, "mean_cost_sd": 0.2}
        ],
        "multiseed_summary": [
            {"regime": "local", "metric": "mean_cost", "mean": 0.5, "sd": 0.1}
        ],
        "client_round_losses": [
            {"regime": "fedavg", "node": "alpha", "round": 1, "training_loss": 0.5}
        ],
        "pre_confusion": [
            {"regime": "local", "actual_class": 0, "predicted_class": 0, "examples": 10}
        ],
    }
    for name in legacy_report.TABLES:
        rows = fixtures.get(name, [{"seed": 100, "mean_cost": 0.5}])
        with (tmp_path / f"{name}.csv").open(
            "w", newline="", encoding="utf-8"
        ) as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
    narrative = "# Diagnóstico histórico\n\nConclusión original.\n\n| field | value |\n| --- | --- |\n| config | base |\n"
    (tmp_path / "resultados.md").write_text(narrative, encoding="utf-8")
    for name in legacy_report.FIGURES:
        (tmp_path / f"{name}.png").write_bytes(b"old")
    (tmp_path / "unrelated.png").write_bytes(b"preserve")
    monkeypatch.setattr(legacy_report, "matplotlib_jpg", lambda *_: _small_jpeg())
    legacy_report.refresh_legacy_report(tmp_path)
    legacy_report.refresh_legacy_report(tmp_path)
    text = (tmp_path / "resultados.md").read_text(encoding="utf-8")
    assert text.count("Conclusión original.") == 1
    assert text.count("## Descargas del diagnóstico histórico") == 1
    assert text.count("| config | base |") == 1
    assert "metadatos_historicos_1.xlsx" in text
    assert len(list(tmp_path.glob("*.xlsx"))) == 13
    assert list(tmp_path.glob("*.png")) == [tmp_path / "unrelated.png"]


def _small_jpeg():
    target = BytesIO()
    Image.new("RGB", (10, 10), "white").save(target, "JPEG", dpi=(600, 600))
    return target.getvalue()
