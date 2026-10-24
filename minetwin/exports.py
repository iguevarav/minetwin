"""Shared, reproducible exports for the UI and offline reports."""

import base64
import html
import json
import math
import re
import unicodedata
from dataclasses import dataclass, replace
from io import BytesIO

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

DPI = 600
CSS_DPI = 96
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class Explanation:
    title: str
    measure: str
    reading: str
    method: str
    limitation: str
    source: str = "SCANIA Component X · artefactos reproducibles de MineTwin"
    result: str = ""

    def items(self) -> list[tuple[str, str]]:
        return [
            ("Título", self.title),
            ("Qué mide", self.measure),
            ("Cómo interpretarlo", self.reading),
            ("Cómo se obtiene", self.method),
            ("Resultado observado", self.result or "Sin observaciones disponibles."),
            ("Límite de la conclusión", self.limitation),
            ("Fuente y contexto", self.source),
        ]

    def markdown(self) -> str:
        return "  \n".join(f"**{label}:** {value}" for label, value in self.items())


def filename(title: str) -> str:
    ascii_title = (
        unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    )
    return re.sub(r"[^a-zA-Z0-9]+", "_", ascii_title).strip("_").lower() or "resultado"


def table_result(rows: list[dict]) -> str:
    if not rows:
        return "Sin filas disponibles; no permite extraer una conclusión numérica."
    # Report only values actually present; never infer causal or clinical conclusions.
    fields = list(dict.fromkeys(key for row in rows for key in row))
    ranges = []
    for key in fields:
        values = [row.get(key) for row in rows]
        numbers = []
        for value in values:
            if value is None or isinstance(value, (bool, np.bool_)):
                continue
            try:
                number = float(value)
            except (TypeError, ValueError):
                continue
            if math.isfinite(number):
                numbers.append(number)
        if numbers and len(ranges) < 4:
            low, high = min(numbers), max(numbers)
            ranges.append(
                f"{key}: {low:.6g}" + (f" a {high:.6g}" if low != high else "")
            )
    return f"{len(rows):,} filas. " + "; ".join(ranges)


def explained_rows(explanation: Explanation, rows: list[dict]) -> Explanation:
    return replace(explanation, result=table_result(rows))


def _cell(value, parse_numbers: bool = False):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, (list, tuple, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if parse_numbers and isinstance(value, str):
        if value in {"True", "False"}:
            return value == "True"
        if re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?", value):
            # Excel stores at most 15 significant digits; retain long IDs as text.
            if re.fullmatch(r"-?\d+", value):
                return int(value) if len(value.lstrip("-")) <= 15 else value
            number = float(value)
            return number if math.isfinite(number) else None
    return value


def table_xlsx(
    rows: list[dict], explanation: Explanation, *, parse_numbers: bool = False
) -> bytes:
    """Preserve numeric cells, missing values and literal text, including '=...'."""
    workbook = Workbook()
    data = workbook.active
    data.title = "Datos"
    fields = list(dict.fromkeys(key for row in rows for key in row))
    data.append([str(key) for key in fields])
    for row in rows:
        data.append(
            [
                _cell(
                    row.get(key),
                    parse_numbers
                    and str(key).lower()
                    not in {
                        "vehicle_id",
                        "id",
                        "group",
                        "feature",
                        "sensor",
                        "variable",
                        "grupo",
                    },
                )
                for key in fields
            ]
        )
    notes = workbook.create_sheet("Interpretación")
    for label, value in explanation.items():
        notes.append([label, value])
    for sheet in (data, notes):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="214D65")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                max(width + 2, 14), 65
            )
            for cell in column:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if sheet is data and fields:
            sheet.auto_filter.ref = sheet.dimensions
    notes.column_dimensions["A"].width = 26
    notes.column_dimensions["B"].width = 110
    for row in notes:
        notes.row_dimensions[row[0].row].height = max(
            30, 15 * math.ceil(len(str(row[1].value)) / 95)
        )
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def jpeg_with_explanation(image: Image.Image, explanation: Explanation) -> bytes:
    """Append a legible explanation and write true 600-DPI JPEG metadata."""
    image = image.convert("RGB")
    width = image.width
    font_size = max(18, round(width / 72))
    for font_name in ("arial.ttf", "DejaVuSans.ttf", "LiberationSans-Regular.ttf"):
        try:
            font = ImageFont.truetype(font_name, size=font_size)
            break
        except OSError:
            continue
    else:
        raise RuntimeError(
            "Instale una fuente TrueType Arial, DejaVu Sans o Liberation Sans para exportar texto en español."
        )
    padding = round(width / 40)
    line_height = round(font.size * 1.45)
    draw = ImageDraw.Draw(image)
    lines = []
    for label, value in explanation.items():
        words = f"{label}: {value}".split()
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and draw.textlength(candidate, font=font) > width - 2 * padding:
                lines.append(current)
                current = word
            else:
                current = candidate
        lines.append(current)
    canvas = Image.new(
        "RGB", (width, image.height + 2 * padding + line_height * len(lines)), "white"
    )
    canvas.paste(image, (0, 0))
    draw = ImageDraw.Draw(canvas)
    for index, line in enumerate(lines):
        draw.text(
            (padding, image.height + padding + index * line_height),
            line,
            font=font,
            fill="#243444",
        )
    target = BytesIO()
    canvas.save(target, "JPEG", dpi=(DPI, DPI), quality=95, subsampling=0)
    return target.getvalue()


def plotly_jpg(figure, explanation: Explanation) -> bytes:
    import plotly.graph_objects as go

    exported = go.Figure(figure)
    exported.update_layout(
        template="plotly_white",
        paper_bgcolor="white",
        plot_bgcolor="white",
        width=768,
        height=max(420, figure.layout.height or 480),
        margin={"l": 90, "r": 45, "t": 65, "b": 85},
        title={"text": explanation.title, "font": {"size": 18}},
    )
    # 768 CSS px / 96 = 8 inches; rasterize natively at 4800 px = 600 DPI.
    raster = exported.to_image(format="png", scale=DPI / CSS_DPI)
    with Image.open(BytesIO(raster)) as image:
        return jpeg_with_explanation(image, explanation)


def matplotlib_jpg(figure, explanation: Explanation) -> bytes:
    figure.tight_layout()
    raster = BytesIO()
    figure.savefig(raster, format="png", dpi=DPI, facecolor="white")
    raster.seek(0)
    with Image.open(raster) as image:
        return jpeg_with_explanation(image, explanation)


def download_link(data: bytes, name: str, mime: str, label: str) -> str:
    encoded = base64.b64encode(data).decode("ascii")
    return (
        f'<a download="{html.escape(name, quote=True)}" '
        f'href="data:{mime};base64,{encoded}">{html.escape(label)}</a>'
    )


def explanation_html(explanation: Explanation) -> str:
    return "".join(
        f"<p><strong>{html.escape(label)}:</strong> {html.escape(value)}</p>"
        for label, value in explanation.items()
    )
